from sqlalchemy.ext.asyncio import AsyncSession
import json
import asyncio
from app.utils import database, prompts, cache, bing, images, gpt, cos
from app.ads import ads
from app.cases import manage_cases
from app.proposal import company

import tempfile

from docx import Document
from docx.shared import Inches


async def generate_remaining_proposal(
    db: AsyncSession, gpt_client, current_user: database.UserLogin
):
    """生成完整提案（技术点和案例可能需要后台生成）"""
    tech_names = cache.get_proposal_techs(current_user.id)
    case_ids = cache.get_proposal_cases(current_user.id)

    proposal = cache.CachedProposal(
        **json.loads(cache.get_record_in_cache(current_user.id))
    )

    async def _generate_tech(tech_name: str):
        if not cache.is_proposal_tech_text_in_cache(current_user.id, tech_name):
            ad = await ads.search_ad_by_tech_point(
                db, proposal.business, proposal.category, tech_name, current_user
            )

            prompt = await prompts.plan_tech_prompt(
                ad=ad,
                tech=tech_name,
                business=proposal.business,
                category=proposal.category,
                user=current_user,
                db=db,
                background=proposal.requirement,
            )
            txt = await gpt.get_gpt_txt_result(prompt, gpt_client)
            cache.add_proposal_tech_text_in_cache(current_user.id, tech_name, txt)
        if not cache.is_proposal_tech_image_in_cache(current_user.id, tech_name):
            image_url = (
                cache.get_ad_in_cache(f"{current_user.id}-{tech_name}")
                or await images.search_image_urls(tech_name, gpt_client)
                or await bing.get_img(tech_name + " 示意图")
            )
            cache.add_proposal_tech_image_in_cache(
                current_user.id, tech_name, image_url
            )

    tech_tasks = [_generate_tech(tech_name) for tech_name in tech_names]

    async def _generate_case(case_id: int):
        if not cache.is_proposal_case_text_in_cache(current_user.id, case_id):
            prompt = await manage_cases.get_case_prompt(
                db=db,
                account_type=current_user.account_type,
                case_id=case_id,
                business=proposal.business,
                category=proposal.category,
                background=proposal.requirement,
            )
            txt = await gpt.get_gpt_txt_result(prompt, gpt_client)
            cache.add_proposal_case_text_in_cache(current_user.id, case_id, txt)

    case_tasks = [_generate_case(case_id) for case_id in case_ids]

    tasks = tech_tasks + case_tasks
    await asyncio.gather(*tasks)


async def generate_proposal(
    db: AsyncSession, gpt_client, current_user: database.UserLogin
):
    """下载完整提案"""
    await generate_remaining_proposal(db, gpt_client, current_user)
    # 获取公司信息
    company_info = await company.get_company_info(db, current_user)
    proposal = cache.CachedProposal(
        **json.loads(cache.get_record_in_cache(current_user.id))
    )
    # 生成word
    document = Document()
    document.add_heading(f"环保{proposal.business}节能减排方案", 0)
    document.add_heading("概述", level=1)
    try:
        document.add_picture(
            await images.get_image_bytes_by_url(
                cache.get_proposal_intro_image_in_cache(current_user.id)
            ),
            width=Inches(4),
        )
    except Exception as e:
        print(e)
    document.add_paragraph(cache.get_proposal_intro_text_in_cache(current_user.id))

    document.add_heading("政策环境", level=1)

    try:
        document.add_picture(
            await images.get_image_bytes_by_url(
                cache.get_proposal_policy_image_in_cache(current_user.id)
            ),
            width=Inches(4),
        )
    except Exception as e:
        print(e)

    document.add_heading("国家政策", level=2)
    document.add_paragraph(cache.get_proposal_policy_nation_in_cache(current_user.id))
    document.add_heading("地方政策", level=2)
    document.add_paragraph(cache.get_proposal_policy_local_in_cache(current_user.id))

    document.add_heading("公司介绍", level=1)

    try:
        document.add_picture(
            await images.get_image_bytes_by_url(company_info.image), width=Inches(4)
        )
    except Exception as e:
        print(e)

    document.add_paragraph(company_info.intro)

    document.add_heading("技术方案", level=1)
    document.add_paragraph(cache.get_proposal_tech_plan_in_cache(current_user.id))

    tech_names = cache.get_proposal_techs(current_user.id)
    for tech_name in tech_names:
        document.add_heading(tech_name, level=2)
        try:
            document.add_picture(
                await images.get_image_bytes_by_url(
                    cache.get_proposal_tech_image_in_cache(current_user.id, tech_name)
                ),
                width=Inches(4),
            )
        except Exception as e:
            print(e)
        document.add_paragraph(
            cache.get_proposal_tech_text_in_cache(current_user.id, tech_name)
        )

    document.add_heading("应用案例", level=1)
    case_ids = cache.get_proposal_cases(current_user.id)
    for case_id in case_ids:
        case_title = cache.get_proposal_case_title_in_cache(current_user.id, case_id)
        document.add_heading(case_title, level=2)
        try:
            document.add_picture(
                await images.get_image_bytes_by_url(
                    cache.get_proposal_case_image_in_cache(current_user.id, case_id)
                ),
                width=Inches(4),
            )
        except Exception as e:
            print(e)
        document.add_paragraph(
            cache.get_proposal_case_text_in_cache(current_user.id, case_id)
        )

    document.save(f"{proposal.record_id}.docx")
    # word上传到cos
    with tempfile.NamedTemporaryFile() as tmp:
        document.save(tmp.name)
        tmp.flush()
        file_path = await cos.upload_file_to_cos(
            tmp.name, f"{proposal.record_id}.docx", "proposals"
        )
    return file_path
